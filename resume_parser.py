"""
resume_parser.py - Robust PDF, DOCX, and Text Resume Parser
Extracts contact information, sections, experience, education, and projects.
"""

import re
import os
from typing import Dict, Any, List

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import docx
except ImportError:
    docx = None


SECTION_HEADERS = {
    "summary": ["summary", "professional summary", "executive summary", "profile", "objective", "career objective", "about me"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "tech stack", "tools & technologies", "skills & tools"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history", "internships"],
    "projects": ["projects", "personal projects", "academic projects", "key projects", "notable projects"],
    "education": ["education", "academic background", "academic history", "qualifications"],
    "certifications": ["certifications", "licenses & certifications", "certificates", "courses", "achievements", "awards", "honors"]
}


def extract_text_from_pdf(filepath: str) -> str:
    """Extract text from a PDF file using PyPDF2."""
    if not PyPDF2:
        raise ImportError("PyPDF2 is not installed.")
    text = ""
    with open(filepath, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text.strip()


def extract_text_from_docx(filepath: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    if not docx:
        raise ImportError("python-docx is not installed.")
    doc = docx.Document(filepath)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text).strip()


def extract_text(filepath: str) -> str:
    """Detect file type and extract text."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(filepath)
    elif ext in [".txt", ".md"]:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def extract_contact_info(text: str) -> Dict[str, Any]:
    """Extract candidate name, email, phone, and online profiles."""
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    first_few_lines = lines[:8] if lines else []
    
    # Email
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    email = email_match.group(0) if email_match else None

    # Phone
    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
    phone = phone_match.group(0) if phone_match else None

    # Links
    linkedin_match = re.search(r'(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    linkedin = f"https://linkedin.com/in/{linkedin_match.group(1)}" if linkedin_match else None

    github_match = re.search(r'(?:https?:\/\/)?(?:www\.)?github\.com\/([a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    github = f"https://github.com/github.com/{github_match.group(1)}" if github_match else None
    
    portfolio_match = re.search(r'(?:portfolio|website):\s*(https?:\/\/[^\s]+)', text, re.IGNORECASE)
    portfolio = portfolio_match.group(1) if portfolio_match else None

    # Name heuristic: First non-empty line that doesn't contain email, phone, or symbols
    name = "Candidate"
    for line in first_few_lines:
        clean = line.strip()
        if (email and email in clean) or (phone and phone in clean) or "@" in clean or "http" in clean or "|" in clean:
            # check if name is before the pipe
            if "|" in clean:
                candidate = clean.split("|")[0].strip()
                if len(candidate) > 2 and len(candidate.split()) <= 4 and not re.search(r'\d', candidate):
                    name = candidate
                    break
            continue
        # Avoid common section titles
        lower_line = clean.lower()
        if any(h in lower_line for h in ["resume", "curriculum vitae", "cv", "developer", "engineer", "profile", "summary"]):
            continue
        if 2 <= len(clean.split()) <= 4 and re.match(r"^[A-Za-z\s.'-]+$", clean):
            name = clean
            break

    return {
        "name": name,
        "email": email or "Not detected",
        "phone": phone or "Not detected",
        "linkedin": linkedin or "Not detected",
        "github": github or "Not detected",
        "portfolio": portfolio or "Not detected"
    }


def parse_sections(text: str) -> Dict[str, str]:
    """Split resume into standard sections using header detection."""
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {
        "header": [],
        "summary": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "education": [],
        "certifications": [],
        "other": []
    }

    current_section = "header"

    def match_header(line: str) -> str:
        clean = line.strip().lower()
        # Clean markdown symbols, colons, dashes
        clean = re.sub(r'^[#*\-_\s]+', '', clean)
        clean = re.sub(r'[:\-–—\s]+$', '', clean).strip()
        
        if len(clean) > 40:
            return None

        for section_name, headers in SECTION_HEADERS.items():
            for h in headers:
                if clean == h or clean == f"{h}:" or clean == f"my {h}":
                    return section_name
        return None

    for line in lines:
        matched = match_header(line)
        if matched:
            current_section = matched
            continue
        sections[current_section].append(line)

    return {sec: "\n".join(content).strip() for sec, content in sections.items()}


def extract_bullet_points(section_text: str) -> List[str]:
    """Extract individual bullet points or meaningful sentences from a section."""
    bullets = []
    lines = section_text.split("\n")
    current_bullet = ""

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current_bullet:
                bullets.append(current_bullet.strip())
                current_bullet = ""
            continue
        
        # Starts with bullet marker (-, *, •, \u2022) or number
        if re.match(r'^[-*•\u2022]\s+', stripped) or re.match(r'^\d+\.\s+', stripped):
            if current_bullet:
                bullets.append(current_bullet.strip())
            current_bullet = re.sub(r'^[-*•\u2022\d.]\s*', '', stripped)
        else:
            if current_bullet:
                current_bullet += " " + stripped
            else:
                if len(stripped) > 20:
                    current_bullet = stripped
                else:
                    bullets.append(stripped)

    if current_bullet:
        bullets.append(current_bullet.strip())

    return [b for b in bullets if len(b) > 10]


def parse_resume(text: str) -> Dict[str, Any]:
    """Complete resume parse returning contacts, structured sections, and bullets."""
    contacts = extract_contact_info(text)
    sections = parse_sections(text)
    
    experience_bullets = extract_bullet_points(sections.get("experience", ""))
    project_bullets = extract_bullet_points(sections.get("projects", ""))

    return {
        "contact": contacts,
        "sections": sections,
        "experience_bullets": experience_bullets,
        "project_bullets": project_bullets,
        "raw_text": text,
        "word_count": len(text.split()),
        "character_count": len(text)
    }
